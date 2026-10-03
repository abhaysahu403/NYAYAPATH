"""Document processing pipeline.

    validate -> store -> job -> extract text (per page) -> OCR pages that have no text layer
    -> clean -> chunk (page + heading preserved) -> embed -> INDEXED

Runs synchronously today. Every attempt is recorded as a DocumentProcessingJob, so a Celery/RQ worker can call
`process_document(doc_id)` later without changing the pipeline.
"""
import hashlib
import io
import logging
import zipfile
from pathlib import Path
from typing import List, Optional, Tuple

from django.conf import settings
from django.utils import timezone

from core.text import chunk_pages
from .models import Document, DocumentChunk, DocumentProcessingJob, DocumentStatus, ProcessingJobStatus
from .storage import absolute_path

logger = logging.getLogger("nyayapath.documents.processing")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EXT_TO_MIME = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".docx": DOCX_MIME}
ALLOWED_EXTENSIONS = set(EXT_TO_MIME)
ALLOWED_MIME_TYPES = set(EXT_TO_MIME.values())
_EXECUTABLE_SIGNATURES = (b"MZ", b"\x7fELF", b"#!", b"\xca\xfe\xba\xbe", b"\xfe\xed\xfa")
OCR_MIN_CHARS_PER_PAGE = 30
OCR_MAX_PAGES = getattr(settings, "OCR_MAX_PAGES", 40)

Pages = List[Tuple[Optional[int], str]]


# ------------------------------------------------------------------ validation
def validate_upload(file) -> str:
    """Validate extension, size and real content (magic bytes). Returns the verified MIME type or raises ValueError."""
    ext = Path(file.name or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"File type '{ext or 'unknown'}' is not allowed. Accepted: PDF, PNG, JPG, JPEG, DOCX.")
    if file.size == 0:
        raise ValueError("The uploaded file is empty.")
    if file.size > settings.MAX_UPLOAD_SIZE_BYTES:
        raise ValueError(f"File is too large. Maximum size is {settings.MAX_UPLOAD_SIZE_MB} MB.")
    file.seek(0)
    head = file.read(16)
    file.seek(0)
    if head.startswith(_EXECUTABLE_SIGNATURES):
        raise ValueError("Executable files are not allowed.")
    ok = {
        ".pdf": head.startswith(b"%PDF"),
        ".png": head.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": head.startswith(b"\xff\xd8\xff"),
        ".jpeg": head.startswith(b"\xff\xd8\xff"),
        ".docx": head.startswith(b"PK\x03\x04"),
    }[ext]
    if not ok:
        raise ValueError("File content does not match its extension.")
    if ext == ".docx":
        try:
            with zipfile.ZipFile(file) as z:
                if "word/document.xml" not in z.namelist():
                    raise ValueError("File is not a valid DOCX document.")
        except zipfile.BadZipFile:
            raise ValueError("File is not a valid DOCX document.")
        finally:
            file.seek(0)
    return EXT_TO_MIME[ext]


def compute_checksum(file) -> str:
    h = hashlib.sha256()
    file.seek(0)
    for block in iter(lambda: file.read(65536), b""):
        h.update(block)
    file.seek(0)
    return h.hexdigest()


# ------------------------------------------------------------------ OCR
def ocr_available() -> bool:
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:  # noqa: BLE001
        return False


def _ocr_langs() -> str:
    """Use configured languages but only those actually installed (so a missing 'hin' pack degrades to English)."""
    import pytesseract
    wanted = [l for l in settings.OCR_LANGUAGES.split("+") if l]
    try:
        have = set(pytesseract.get_languages(config=""))
        wanted = [l for l in wanted if l in have] or ["eng"]
    except Exception:  # noqa: BLE001
        pass
    return "+".join(wanted)


def ocr_image(image) -> str:
    import pytesseract
    return pytesseract.image_to_string(image, lang=_ocr_langs()) or ""


def _render_pdf_page(path: str, index: int):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(path)
    try:
        return pdf[index].render(scale=200 / 72).to_pil()
    finally:
        pdf.close()


# ------------------------------------------------------------------ extraction
def extract_pdf(path: str) -> Tuple[Pages, int, bool]:
    """Returns (pages, page_count, ocr_used). Pages without a text layer are OCR'd when OCR is available."""
    import pypdf
    reader = pypdf.PdfReader(path)
    page_count = len(reader.pages)
    pages: Pages = []
    ocr_used = False
    can_ocr = ocr_available()
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            text = ""
        if len(text.strip()) < OCR_MIN_CHARS_PER_PAGE and can_ocr and i < OCR_MAX_PAGES:
            try:
                text = ocr_image(_render_pdf_page(path, i)) or text
                ocr_used = True
            except Exception as exc:  # noqa: BLE001
                logger.warning("page_ocr_failed", extra={"service": "ocr", "error_type": type(exc).__name__})
        pages.append((i + 1, text))
    return pages, page_count, ocr_used


def extract_docx(path: str) -> Tuple[Pages, int, bool]:
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return [(None, "\n\n".join(parts))], 1, False


def extract_image(path: str) -> Tuple[Pages, int, bool]:
    from PIL import Image
    with Image.open(path) as img:
        img.load()
        return [(1, ocr_image(img))], 1, True


def extract(doc: Document) -> Tuple[Pages, int, bool]:
    path = str(absolute_path(doc.file_path))
    if doc.mime_type == "application/pdf":
        return extract_pdf(path)
    if doc.mime_type == DOCX_MIME:
        return extract_docx(path)
    if doc.mime_type in {"image/png", "image/jpeg"}:
        return extract_image(path)
    raise ValueError("Unsupported document type")


# ------------------------------------------------------------------ pipeline
def _finish_job(job, status, error_code=None):
    job.status = status
    job.error_code = error_code
    job.finished_at = timezone.now()
    job.save(update_fields=["status", "error_code", "finished_at"])


def process_document(doc_id: str) -> bool:
    """Full pipeline for one document. Returns True when the document is searchable (INDEXED)."""
    try:
        doc = Document.objects.get(pk=doc_id)
    except Document.DoesNotExist:
        return False
    prior = doc.jobs.count()
    job = DocumentProcessingJob.objects.create(document=doc, status=ProcessingJobStatus.RUNNING, attempts=prior + 1,
                                               started_at=timezone.now())
    try:
        doc.status = DocumentStatus.PROCESSING
        doc.processing_error = None
        doc.save(update_fields=["status", "processing_error"])

        pages, page_count, ocr_used = extract(doc)
        full_text = "\n\n".join(t for _, t in pages).strip()
        doc.page_count = page_count
        doc.ocr_performed = ocr_used
        if len(full_text) < 20:
            needs_ocr = doc.mime_type != DOCX_MIME and not ocr_available()
            doc.status = DocumentStatus.OCR_REQUIRED if needs_ocr else DocumentStatus.FAILED
            doc.processing_error = ("OCR is required but not available on this server." if needs_ocr
                                    else "No readable text could be extracted from this document.")
            doc.save(update_fields=["status", "processing_error", "page_count", "ocr_performed"])
            _finish_job(job, ProcessingJobStatus.FAILED, "OCR_UNAVAILABLE" if needs_ocr else "NO_TEXT")
            return False

        doc.extracted_text = full_text[:2_000_000]
        doc.status = DocumentStatus.TEXT_EXTRACTED
        doc.save(update_fields=["extracted_text", "status", "page_count", "ocr_performed"])

        chunks = chunk_pages(pages)
        DocumentChunk.objects.filter(document=doc).delete()
        DocumentChunk.objects.bulk_create([
            DocumentChunk(document=doc, chunk_index=i, text=c.text, page_number=c.page_number, section_title=c.section_title,
                          char_start=c.char_start, char_end=c.char_end,
                          chunk_metadata={"document_id": str(doc.id), "chunk_total": len(chunks), "ocr": ocr_used})
            for i, c in enumerate(chunks)], batch_size=100)
        embed_document_chunks(doc)

        doc.status = DocumentStatus.INDEXED
        doc.processed_at = timezone.now()
        doc.save(update_fields=["status", "processed_at"])
        _finish_job(job, ProcessingJobStatus.SUCCEEDED)
        return True
    except Exception as exc:  # noqa: BLE001 - never leak internals; record a code
        logger.error("process_document_failed", extra={"service": "documents", "error_type": type(exc).__name__})
        doc.status = DocumentStatus.FAILED
        doc.processing_error = "The document could not be processed."
        doc.save(update_fields=["status", "processing_error"])
        _finish_job(job, ProcessingJobStatus.FAILED, type(exc).__name__[:60])
        return False


def embed_document_chunks(doc: Document) -> int:
    """Embed chunks that lack a vector. Failure leaves chunks keyword-searchable and is retried by a command."""
    from ai.embedding import EmbeddingService
    chunks = list(DocumentChunk.objects.filter(document=doc, embedding__isnull=True).order_by("chunk_index"))
    if not chunks:
        return 0
    try:
        vectors = EmbeddingService().embed_documents([c.text for c in chunks])
    except Exception as exc:  # noqa: BLE001
        logger.warning("embed_failed", extra={"service": "documents", "error_type": type(exc).__name__})
        return 0
    for c, v in zip(chunks, vectors):
        c.embedding = v
    DocumentChunk.objects.bulk_update(chunks, ["embedding"], batch_size=50)
    return len(chunks)
